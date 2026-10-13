"""Read-only exported workbook validation; save a separate evidence manifest."""
import hashlib
import json
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/table-management-20261006'
NAMES = {'first': '第1次', 'second': '第2次', 'third': '第3次',
         'fourth': '第4次', 'fifth': '第5次'}


def read(path):
    return json.loads(path.read_text(encoding='utf-8'))


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    source = read(OUT / 'inputs.json')
    for entry in source['inventory']:
        assert digest(ROOT / entry['path']) == entry['sha256'], entry['path']
    for path, expected in source['source_sha256'].items():
        assert digest(ROOT / path) == expected, path
    path = OUT / 'current/project-evaluation.xlsx'
    book = load_workbook(path, read_only=True, data_only=True)
    assert book.sheetnames == ['总体对比', '评分', '生成回复', '重复复核改动', '退款补充清单', '文件索引']
    actual = list(book['评分'].values)[1:]
    expected = [(NAMES[r[0]], *r[1:]) for r in source['scores']]
    # Excel stores empty strings as blank cells.
    normalize = lambda rows: [tuple(None if v == '' else v for v in row) for row in rows]
    assert normalize(actual) == normalize(expected)
    assert len(actual) == 567
    replies = list(book['生成回复'].values)[1:]
    expected_replies = [(NAMES[r[0]], *r[1:4], r[4] if r[2] == 1 else '', r[5], r[6])
                        for r in source['replies']]
    assert normalize(replies) == normalize(expected_replies)
    grouped = defaultdict(list)
    for r in replies:
        grouped[r[0], r[6]].append((r[2], r[3], r[5] or ''))
    assert len(grouped) == 567
    paths = [p for p in source['source_sha256'] if p.endswith('replies.jsonl')]
    for model, p in zip(('second', 'third', 'fourth', 'fifth', 'first'), paths):
        for line in (ROOT / p).read_text(encoding='utf-8').splitlines():
            r = json.loads(line)
            chunks = sorted(grouped[NAMES[model], r['id']])
            assert [c[0] for c in chunks] == list(range(1, len(chunks) + 1))
            assert all(c[1] == len(chunks) for c in chunks)
            assert ''.join(c[2] for c in chunks) == r.get('generated_response', r.get('response'))
    assert normalize(list(book['退款补充清单'].values)[1:]) == normalize(source['supplement'])
    rules = read(ROOT / 'outputs/refund-supplement-20261006-v1.1/approved-rules.json')
    assert source['supplement'] == [[r['review_id'], len(r['members']), r['before'], r['after'],
                                    r['reason'], ', '.join(m['id'].rsplit(':', 1)[-1] for m in r['members'])]
                                   for r in rules]
    assert len(rules) == 172
    assert normalize(list(book['重复复核改动'].values)[1:]) == normalize(
        [(NAMES[r[0]], *r[1:]) for r in source['repetition_changes']])
    summary = book['总体对比']
    for row, model in ((6, 'first'), (7, 'second'), (8, 'third'), (9, 'fourth'), (10, 'fifth')):
        entries = [r for r in source['scores'] if r[0] == model]
        assert summary[f'G{row}'].value == len(entries)
        assert abs(summary[f'F{row}'].value - source['losses'][model]) < 1e-12
        for column, index in [('H', 4), ('I', 5), ('J', 6)]:
            assert abs(summary[f'{column}{row}'].value - sum(r[index] for r in entries) / len(entries)) < 1e-12
        assert summary[f'K{row}'].value == sum(r[8] for r in entries)
        assert summary[f'L{row}'].value == (sum(r[7] for r in entries)
                                                  if model in ('fourth', 'fifth') else 'n.a.')
    assert summary['L9'].value == 58 and summary['L10'].value == 63
    assert summary['F29'].value == 5 and summary['G29'].value == 135
    errors = [(s.title, c.coordinate, c.value) for s in book for row in s for c in row
              if c.data_type == 'e']
    assert not errors, errors
    book.close()
    evidence = {'frozen': True, 'authoritative_current_workbook': 'current/project-evaluation.xlsx',
                'scores_checked': 567, 'reply_chunks_checked': len(replies),
                'full_replies_reconstructed': 567, 'supplement_rules_checked': 172,
                'historical_workbooks_unchanged': len(source['inventory']),
                'cached_excel_errors': errors, 'fourth_repeat': 58, 'fifth_repeat': 63,
                'paired_net_new_repeat': 5, 'formula_recalculation_checked_by_builder': True,
                'sha256': {str(p.relative_to(ROOT)).replace('\\', '/'): digest(p)
                           for p in (path, OUT / 'inputs.json', OUT / 'historical-inventory.json',
                                     OUT / 'build-current.mjs', OUT / 'INDEX.md',
                                     ROOT / 'docs/project-roadmap.md',
                                     ROOT / 'docs/evaluation-consolidation-20261006.md', Path(__file__))},
                'source_sha256': source['source_sha256'],
                'historical_sha256': {r['path']: r['sha256'] for r in source['inventory']},
                'scope_limit': 'AI diagnostic scores; supplement semantic boundary review pending; no training activation'}
    # Evidence only; never modifies a workbook or a source score.
    (OUT / 'validation.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2), encoding='utf-8')
    print({k: v for k, v in evidence.items() if k != 'sha256'})


if __name__ == '__main__':
    main()
