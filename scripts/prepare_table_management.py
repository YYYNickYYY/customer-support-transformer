"""Read-only XLSX inventory and authoritative inputs for one current workbook.

Run with bundled artifact-analysis Python. Never edit historical workbooks.
"""
import hashlib
import json
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/table-management-20261006'
def read(path):return json.loads((ROOT/path).read_text(encoding='utf-8'))
def digest(path):return hashlib.sha256(path.read_bytes()).hexdigest()
def main():
    if OUT.exists():raise ValueError('Table management already exists')
    files=sorted((ROOT/'outputs').rglob('*.xlsx'))
    inventory=[]
    for p in files:
        formulas=load_workbook(p,read_only=True,data_only=False)
        cached=load_workbook(p,read_only=True,data_only=True)
        sheets=[];errors=[];missing_caches=0
        for sheet in formulas:
            count=0
            cache_rows=list(cached[sheet.title].iter_rows())
            for row in sheet:
                for cell in row:
                    if not hasattr(cell,'coordinate'):continue
                    c=cache_rows[cell.row-1][cell.column-1]
                    if cell.data_type=='f':
                        count+=1
                        if c.value is None:missing_caches+=1
                    if c.data_type=='e':errors.append(f'{sheet.title}!{cell.coordinate}: {c.value}')
            sheets.append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'formulas':count})
        formulas.close();cached.close()
        reason='历史评分及实验快照，后续不在此表继续追加。'
        if 'dev-expanded' in str(p):reason='历史135条汇总；第四/第五重复55/66已被58/63替代，其他维度原分保留。'
        elif 'dev-risk-review' not in str(p) and '20261004' in str(p):
            reason='历史27条及因素表；旧风险v1.0不能与v1.1混用，最新来源见统一表。'
        inventory.append({'path':p.relative_to(ROOT).as_posix(),'sha256':digest(p),
            'status':'历史留档，不作为当前汇总','reason':reason,'sheets':sheets,
            'cached_formula_errors':errors,'formulas_without_saved_cache':missing_caches})
    revised=read('outputs/repetition-review-20261006-v1/ai-expert-ratings-repetition-revised.json')
    material=read('outputs/repetition-review-20261006-v1/material.json')
    mapping=read('outputs/repetition-review-20261006-v1/mapping.json')
    repeated=read('outputs/repetition-review-20261006-v1/reviews.json')
    repeats={(r['model'],r['id']):r for r in repeated}
    source=read('outputs/dev-expanded-fifth-20261005/ai-expert-ratings.json')
    first=read('outputs/dev-risk-review-20261004-v1-1/experiment-1-ratings.json')
    rows=[{**r,'model':'first','cohort':'original27'} for r in first['ratings']]+revised['ratings']
    assert len(rows)==567 and len({(r['model'],r['id']) for r in rows})==567
    replies={}
    paths={'second':'outputs/dev-expanded-20261005-v1/second/replies.jsonl',
           'third':'outputs/dev-expanded-20261005-v1/third/replies.jsonl',
           'fourth':'outputs/dev-expanded-20261005-v1/fourth/replies.jsonl',
           'fifth':'outputs/dev-expanded-fifth-20261005/replies.jsonl',
           'first':'runs/20261004T015453-full-22c1eff31be445aa83c00458c4582e6a/dev-checks/20261004T091445-fdf38127e5b2435597a623e0b2ff61c5/replies.jsonl'}
    for model,p in paths.items():
        for line in (ROOT/p).read_text(encoding='utf-8').splitlines():
            r=json.loads(line);replies[model,r['id']]=r
    score_rows=[];reply_rows=[]
    for r in sorted(rows,key=lambda x:(x['model'],x['id'])):
        key=(r['model'],r['id']);reply=replies[key];review=repeats.get(key)
        origin='repetition-review-20261006-v1' if review else ('risk-review-v1.1 first27' if r['model']=='first' else 'fifth-expanded historical snapshot')
        repeat_version='统一复核v1.0' if review else '历史v1.0，未同口径复核'
        score_rows.append([r['model'],r['id'].rsplit(':',1)[-1],reply['intent'],r['cohort'],
            r['intent_score'],r['helpfulness_score'],r['fluency_score'],r['repetition'],
            r['unsupported_claim'],repeat_version,review['old_repetition'] if review else r['repetition'],
            review['reason'] if review else r.get('notes',''),r['id'],origin])
        response=reply.get('generated_response',reply.get('response'))
        assert isinstance(response,str)
        parts=[response[i:i+450] for i in range(0,len(response),450)] or ['']
        assert ''.join(parts)==response
        for i,part in enumerate(parts,1):
            reply_rows.append([r['model'],r['id'].rsplit(':',1)[-1],i,len(parts),reply['instruction'],part,r['id']])
    losses={**revised['best_dev_loss_unchanged'],'first':read('runs/20261004T015453-full-22c1eff31be445aa83c00458c4582e6a/summary.json')['best_dev_loss']}
    changes=read('outputs/refund-supplement-20261006-v1.1/changes-preview.json')
    rules=read('outputs/refund-supplement-20261006-v1.1/approved-rules.json')
    source_paths=['outputs/repetition-review-20261006-v1/ai-expert-ratings-repetition-revised.json',
        'outputs/repetition-review-20261006-v1/review-freeze.json',
        'outputs/dev-expanded-fifth-20261005/ai-expert-ratings.json',
        'outputs/dev-risk-review-20261004-v1-1/experiment-1-ratings.json',
        'outputs/refund-supplement-20261006-v1.1/manifest.json',
        'outputs/refund-supplement-20261006-v1.1/approved-rules.json',*paths.values()]
    inputs={'inventory':inventory,'scores':score_rows,'replies':reply_rows,'losses':losses,
        'repetition_changes':[[r['model'],r['id'].rsplit(':',1)[-1],r['old_repetition'],r['repetition'],
            ' / '.join(r['evidence']),r['reason'],r['id']] for r in repeated if r['changed']],
        'supplement':[[r['review_id'],len(r['members']),r['before'],r['after'],r['reason'],
            ', '.join(m['id'].rsplit(':',1)[-1] for m in r['members'])] for r in rules],
        'supplement_manifest':read('outputs/refund-supplement-20261006-v1.1/manifest.json'),
        'source_sha256':{p:digest(ROOT/p) for p in source_paths},
        'paired_repetition':revised['paired_repetition'],
        'historical_workbooks_mutated':False,'source_ratings':567,
        'reason_for_not_moving':'历史表格被冻结清单/日志原路径引用，保留原位以保持证据链；通过唯一当前入口和索引归档。'}
    OUT.mkdir(parents=True)
    (OUT/'inputs.json').write_text(json.dumps(inputs,ensure_ascii=False,indent=2),encoding='utf-8')
    (OUT/'historical-inventory.json').write_text(json.dumps(inventory,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Historical XLSX',len(inventory),'scores',len(score_rows),'reply text parts',len(reply_rows))
    print('Cached errors',sum(len(r['cached_formula_errors']) for r in inventory))
if __name__=='__main__':main()
